"""CSV / TSV / XLSX parser.

Each sheet becomes one ParsedPage holding a Markdown table. The chunker splits
big tables between rows and repeats the header in every chunk, so a question
about row 5000 still sees the column names.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import logging
from pathlib import Path

from notebook_rag.config import get_settings
from notebook_rag.ingestion.parsers._tables import rows_to_markdown
from notebook_rag.models import ParsedPage

log = logging.getLogger(__name__)


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeError:
            continue
    return raw.decode("cp1252", errors="replace")


def _fmt(value: object) -> object:
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ", timespec="minutes") if (value.hour or value.minute) else value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return value


class SheetParser:
    def parse(self, path: Path, doc_id: str) -> list[ParsedPage]:
        ext = path.suffix.lower()
        if ext in (".xlsx", ".xlsm"):
            return self._parse_xlsx(path, doc_id)
        return self._parse_csv(path, doc_id, tab=(ext == ".tsv"))

    # ── CSV / TSV ─────────────────────────────────────────────────────────────
    def _parse_csv(self, path: Path, doc_id: str, tab: bool) -> list[ParsedPage]:
        max_rows = get_settings().max_sheet_rows
        text = _decode(path.read_bytes())
        if tab:
            dialect: csv.Dialect | type[csv.Dialect] = csv.excel_tab
        else:
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
        rows: list[list[str]] = []
        truncated = False
        for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
            if i > max_rows:
                truncated = True
                break
            rows.append(row)
        return self._page(doc_id, path.stem, rows, truncated, max_rows)

    # ── XLSX ──────────────────────────────────────────────────────────────────
    def _parse_xlsx(self, path: Path, doc_id: str) -> list[ParsedPage]:
        from openpyxl import load_workbook

        max_rows = get_settings().max_sheet_rows
        wb = load_workbook(str(path), read_only=True, data_only=True)
        pages: list[ParsedPage] = []
        try:
            for ws in wb.worksheets:
                rows: list[list[object]] = []
                truncated = False
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > max_rows:
                        truncated = True
                        break
                    rows.append([_fmt(v) for v in row])
                pages += self._page(doc_id, ws.title, rows, truncated, max_rows)
        finally:
            wb.close()
        return pages

    @staticmethod
    def _page(doc_id: str, section: str, rows, truncated: bool, max_rows: int) -> list[ParsedPage]:
        md = rows_to_markdown(rows)
        if not md:
            return []
        meta: dict = {}
        if truncated:
            meta["warning"] = (
                f"Sheet '{section}' has more than {max_rows} rows; only the first {max_rows} rows were indexed."
            )
        return [ParsedPage(doc_id=doc_id, page=None, section=section, text=md, metadata=meta)]
