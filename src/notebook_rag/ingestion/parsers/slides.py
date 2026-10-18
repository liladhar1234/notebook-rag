"""PPTX parser (python-pptx): one ParsedPage per slide (page = slide number)."""
from __future__ import annotations

from pathlib import Path

from notebook_rag.ingestion.parsers._tables import rows_to_markdown
from notebook_rag.models import ParsedPage


def _walk(shapes):
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # group
            yield from _walk(shape.shapes)
        else:
            yield shape


class SlidesParser:
    def parse(self, path: Path, doc_id: str) -> list[ParsedPage]:
        from pptx import Presentation

        prs = Presentation(str(path))
        pages: list[ParsedPage] = []
        for number, slide in enumerate(prs.slides, start=1):
            title = ""
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
            blocks: list[str] = [f"Slide {number}: {title}" if title else f"Slide {number}"]
            for shape in _walk(slide.shapes):
                if getattr(shape, "has_table", False) and shape.has_table:
                    md = rows_to_markdown(
                        [[c.text for c in row.cells] for row in shape.table.rows]
                    )
                    if md:
                        blocks.append(md)
                elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    text = "\n".join(
                        p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                    )
                    if text and text != title:
                        blocks.append(text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    blocks.append(f"Speaker notes: {notes}")
            if len(blocks) > 1:
                pages.append(
                    ParsedPage(
                        doc_id=doc_id, page=number, section=title or None,
                        text="\n\n".join(blocks),
                    )
                )
        return pages
